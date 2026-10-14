from __future__ import annotations

import base64
import hashlib
import mimetypes
import os
import zipfile
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote

from .config import Config
from .exports import decode_text
from .media_crypto import decrypt_media


TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".log", ".html", ".ics"}


def inside(path: Path, root: Path) -> bool:
    return path.resolve().is_relative_to(root.resolve())


def sniff(path: Path) -> str:
    with path.open("rb") as f:
        h = f.read(16)
    return sniff_bytes(h, path.name)


def sniff_bytes(h: bytes, name: str = "") -> str:
    if h.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if h.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if h.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if h.startswith(b"RIFF") and h[8:12] == b"WEBP":
        return "image/webp"
    if h.startswith(b"%PDF-"):
        return "application/pdf"
    # Do not mistake encrypted thumbnail blobs for images based on extension.
    if Path(name).suffix.lower() in (".eimg", ".edb"):
        return "application/x-line-encrypted"
    return mimetypes.guess_type(name)[0] or "application/octet-stream"


class RoomFiles:
    def __init__(self, config: Config):
        self.config = config
        self.index: dict[tuple[str, str], tuple[Path, Path]] = {}
        self.linked: dict[tuple[str, str], dict] = {}

    def register_linked(self, room: str, message_id: str, info: dict, content_type: int) -> list[dict]:
        """Paths must come from a room-scoped message query, and stay in approved roots."""
        self.config.room(room)
        roots = [self.config.line_root / "Cache", *self.roots(room),
                 *[self.config.resolve(p) for p in self.config.raw.get("linked_file_roots", [])]]
        descriptors = []
        original_field = "downloadedPath" if info.get("downloadedPath") else "path"
        for field, variant in ((original_field, "original"), ("thumbPath", "thumbnail")):
            value = info.get(field)
            if not isinstance(value, str) or not value:
                continue
            if value.startswith(("\\\\", "//")) or not Path(value).is_absolute() or "\0" in value:
                continue
            path = Path(value).resolve()
            root = next((r for r in roots if inside(path, r)), None)
            if root is None:
                continue
            file_id = "msg:" + message_id + ":" + variant
            filename = unquote(str(info.get("fileName") or path.name))
            key = info.get("keyMaterial")
            exists = path.is_file()
            encrypted = exists and sniff(path) == "application/x-line-encrypted"
            descriptor = {"file_id": file_id, "message_id": message_id, "name": filename,
                          "variant": variant, "exists": exists,
                          "size": path.stat().st_size if exists else None,
                          "encrypted": encrypted, "key_available": bool(key),
                          "readable": exists and (not encrypted or bool(key)),
                          "mime_type": mimetypes.guess_type(filename)[0] or "application/octet-stream",
                          "source": "room_scoped_database_reference"}
            self.linked[room, file_id] = {"path": path, "root": root, "key": key,
                                         "type": content_type, "descriptor": descriptor,
                                         "expected_size": info.get("downloadedFileSize") if field == "downloadedPath" else None}
            descriptors.append(descriptor)
        if not descriptors and content_type in (1, 2, 3, 14):
            descriptors.append({"file_id": "msg:" + message_id + ":original", "message_id": message_id,
                                "name": unquote(str(info.get("fileName") or "Unavailable media")),
                                "variant": "original", "exists": False, "readable": False,
                                "source": "room_scoped_database_reference",
                                "reason": "No local path inside the configured linked_file_roots"})
        return descriptors

    def roots(self, room: str) -> list[Path]:
        return self.config.paths(room, "attachment_roots") + self.config.paths(room, "cache_roots")

    def scan(self, room: str, query: str = "", limit: int = 100, offset: int = 0) -> dict:
        self.config.room(room)
        items = []
        visited = 0
        capped = False
        errors = []
        for root in self.roots(room):
            if not root.is_dir():
                continue
            def onerror(exc):
                errors.append(str(exc))
            for directory, dirs, names in os.walk(root, followlinks=False, onerror=onerror):
                # Also excludes Windows junctions resolving outside this room root.
                dirs[:] = sorted(d for d in dirs if inside(Path(directory) / d, root))
                for name in sorted(names):
                    visited += 1
                    if visited > self.config.max_scan:
                        capped = True
                        break
                    path = Path(directory) / name
                    if not inside(path, root):
                        continue
                    relative = path.relative_to(root).as_posix()
                    if query.casefold() not in relative.casefold():
                        continue
                    file_id = hashlib.sha256((room + "\0" + str(root) + "\0" + relative).encode()).hexdigest()[:24]
                    try:
                        stat = path.stat()
                        mime = sniff(path)
                    except OSError as exc:
                        errors.append(str(exc))
                        continue
                    self.index[room, file_id] = (path, root)
                    items.append({"file_id": file_id, "name": name, "relative_path": relative,
                                  "size": stat.st_size,
                                  "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                                  "mime_type": mime,
                                  "readable": mime != "application/x-line-encrypted"})
                if capped:
                    break
            if capped:
                break
        items.sort(key=lambda r: (r["modified_at"], r["file_id"]), reverse=True)
        return {"files": items[offset:offset + limit], "matched": len(items),
                "has_more": offset + limit < len(items), "scan_capped": capped,
                "errors": errors[:10], "scope": "explicitly configured room directories"}

    def locate(self, room: str, file_id: str) -> Path:
        self.config.room(room)
        linked = self.linked.get((room, file_id))
        if linked:
            path = linked["path"]
            if not inside(path, linked["root"]) or not path.is_file():
                raise ValueError("Referenced LINE attachment is unavailable")
            if path.stat().st_size > self.config.max_file:
                raise ValueError("File exceeds configured max_file_bytes")
            if linked.get("expected_size") is not None and path.stat().st_size != int(linked["expected_size"]):
                raise ValueError("Downloaded file size changed since LINE recorded it")
            return path
        if (room, file_id) not in self.index:
            self.scan(room, limit=1)
        item = self.index.get((room, file_id))
        if not item:
            raise ValueError("Unknown room/file_id; use list_attachments first")
        path, root = item
        if not inside(path, root) or root not in self.roots(room) or not path.is_file():
            raise ValueError("File is no longer inside the configured room directory")
        if path.stat().st_size > self.config.max_file:
            raise ValueError("File exceeds configured max_file_bytes")
        return path

    def content(self, room: str, file_id: str) -> tuple[bytes, str, str]:
        path = self.locate(room, file_id)
        with path.open("rb") as source:
            data = source.read(self.config.max_file + 1)
        if len(data) > self.config.max_file:
            raise ValueError("File grew beyond configured max_file_bytes")
        linked = self.linked.get((room, file_id))
        name = linked["descriptor"]["name"] if linked else path.name
        if sniff(path) == "application/x-line-encrypted":
            if not linked or not linked["key"]:
                raise ValueError("Encrypted cache has no room-linked media key")
            data = decrypt_media(data, linked["key"], video=linked["type"] == 2 and linked["descriptor"]["variant"] == "original")
        return data, name, sniff_bytes(data[:16], name)

    def read_bytes(self, room: str, file_id: str, offset: int, length: int) -> dict:
        if offset < 0 or not 1 <= length <= 262144:
            raise ValueError("offset >= 0; length must be 1..262144")
        path = self.locate(room, file_id)
        content, name, mime = self.content(room, file_id)
        data = content[offset:offset + length]
        return {"file_id": file_id, "name": name, "mime_type": mime, "encoding": "base64",
                "offset": offset, "data": base64.b64encode(data).decode(),
                "next_offset": offset + len(data), "eof": offset + len(data) >= len(content)}

    def read_text(self, room: str, file_id: str, offset: int, max_chars: int) -> dict:
        if offset < 0 or not 1 <= max_chars <= self.config.max_chars:
            raise ValueError("Invalid character offset/max_chars")
        path = self.locate(room, file_id)
        data, name, mime = self.content(room, file_id)
        ext = Path(name).suffix.lower()
        budget = offset + max_chars + 1
        if budget > 2_000_000:
            raise ValueError("Text extraction offset exceeds 2,000,000 characters")
        if ext in TEXT_EXTENSIONS:
            text = decode_text(data)
        elif mime == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDF is unsupported")
            parts, count = [], 0
            for number, page in enumerate(reader.pages, 1):
                part = f"\n[Page {number}]\n" + (page.extract_text() or "")
                parts.append(part)
                count += len(part)
                if count >= budget:
                    break
            text = "".join(parts)
        elif ext in (".docx", ".xlsx"):
            with zipfile.ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(e.file_size for e in entries) > 200 * 1024 * 1024:
                    raise ValueError("Office archive is too large after decompression")
            parts, count = [], 0
            if ext == ".docx":
                from docx import Document
                document = Document(BytesIO(data))
                for p in document.paragraphs:
                    parts.append(p.text)
                    count += len(p.text) + 1
                    if count >= budget:
                        break
                if count < budget:
                    for table in document.tables:
                        for row in table.rows:
                            part = "\t".join(cell.text for cell in row.cells)
                            parts.append(part)
                            count += len(part) + 1
                            if count >= budget:
                                break
                        if count >= budget:
                            break
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    rows_seen = 0
                    for sheet in workbook:
                        parts.append(f"[Sheet: {sheet.title}]")
                        for row in sheet.iter_rows(values_only=True):
                            rows_seen += 1
                            if rows_seen > 100000:
                                raise ValueError("Workbook row limit exceeded")
                            part = "\t".join("" if v is None else str(v) for v in row)
                            parts.append(part)
                            count += len(part) + 1
                            if count >= budget:
                                break
                        if count >= budget:
                            break
                finally:
                    workbook.close()
            text = "\n".join(parts)
        else:
            raise ValueError("No text extractor for this format; images use read_image, other files use read_file_bytes")
        result = text[offset:offset + max_chars]
        return {"file_id": file_id, "name": name, "text": result,
                "offset": offset, "next_offset": offset + len(result),
                "has_more": len(text) > offset + max_chars,
                "note": "Scanned PDFs/images require OCR; audio/video transcription is not implemented"}
